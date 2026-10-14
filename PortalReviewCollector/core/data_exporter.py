"""
Data Exporter Module
Exports scraped reviews to styled multi-sheet Excel (.xlsx), CSV, and JSON formats.
"""
import os
import json
import csv
from datetime import datetime
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .analyzer import ReviewAnalyzer


class DataExporter:
    @classmethod
    def export_excel(cls, reviews: List[Dict[str, Any]], filepath: str, title: str = "3대 포털 리뷰 수집 리포트") -> str:
        """Exports review list and analysis summary into a formatted Excel file."""
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
        wb = openpyxl.Workbook()

        # Styles
        font_title = Font(name="맑은 고딕", size=16, bold=True, color="1E293B")
        font_sub = Font(name="맑은 고딕", size=10, italic=True, color="64748B")
        font_header = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        font_body = Font(name="맑은 고딕", size=10, color="0F172A")
        font_bold = Font(name="맑은 고딕", size=10, bold=True, color="0F172A")

        fill_primary = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
        fill_naver = PatternFill(start_color="03C75A", end_color="03C75A", fill_type="solid") # Naver Green
        fill_kakao = PatternFill(start_color="FEE500", end_color="FEE500", fill_type="solid") # Kakao Yellow
        fill_google = PatternFill(start_color="4285F4", end_color="4285F4", fill_type="solid") # Google Blue
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_accent = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        analysis = ReviewAnalyzer.analyze(reviews)

        # ---------------- Sheet 1: Dashboard Summary ----------------
        ws_dash = wb.active
        ws_dash.title = "📊 수집요약 및 통계"
        ws_dash.views.sheetView[0].showGridLines = True

        ws_dash["A1"] = title
        ws_dash["A1"].font = font_title
        ws_dash["A2"] = f"리포트 생성 일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | 총 수집 리뷰: {len(reviews)}건"
        ws_dash["A2"].font = font_sub

        # Table 1: Portal Summary
        ws_dash["A4"] = "포털별 수집 현황 및 평균 평점"
        ws_dash["A4"].font = font_bold

        headers_p = ["포털", "수집 리뷰 수", "수집 비중(%)", "평균 평점 (5점 만점)"]
        for col_idx, h in enumerate(headers_p, 1):
            cell = ws_dash.cell(row=5, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_primary
            cell.alignment = align_center
            cell.border = thin_border

        curr_row = 6
        for p_name in ["Naver", "Kakao", "Google"]:
            cnt = analysis["portal_breakdown"].get(p_name, 0)
            ratio = round((cnt / len(reviews) * 100), 1) if len(reviews) > 0 else 0
            avg_rate = analysis["average_ratings"].get(p_name, "-")

            ws_dash.cell(row=curr_row, column=1, value=p_name).alignment = align_center
            ws_dash.cell(row=curr_row, column=2, value=cnt).alignment = align_right
            ws_dash.cell(row=curr_row, column=3, value=f"{ratio}%").alignment = align_right
            ws_dash.cell(row=curr_row, column=4, value=avg_rate).alignment = align_right

            for col_idx in range(1, 5):
                c = ws_dash.cell(row=curr_row, column=col_idx)
                c.font = font_body
                c.border = thin_border
            curr_row += 1

        # Overall summary row
        ws_dash.cell(row=curr_row, column=1, value="전체 합계").alignment = align_center
        ws_dash.cell(row=curr_row, column=2, value=len(reviews)).alignment = align_right
        ws_dash.cell(row=curr_row, column=3, value="100.0%").alignment = align_right
        ws_dash.cell(row=curr_row, column=4, value=analysis["overall_average"]).alignment = align_right
        for col_idx in range(1, 5):
            c = ws_dash.cell(row=curr_row, column=col_idx)
            c.font = font_bold
            c.fill = fill_accent
            c.border = thin_border

        # Table 2: Rating Distribution
        r_dist_row = curr_row + 2
        ws_dash.cell(row=r_dist_row, column=1, value="별점 점수대별 분포").font = font_bold
        headers_r = ["별점 구간", "리뷰 수", "비율(%)"]
        for col_idx, h in enumerate(headers_r, 1):
            cell = ws_dash.cell(row=r_dist_row + 1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_primary
            cell.alignment = align_center
            cell.border = thin_border

        c_row = r_dist_row + 2
        for star in [5, 4, 3, 2, 1, "none"]:
            label = f"{star}점" if star != "none" else "별점 미기재"
            s_cnt = analysis["rating_distribution"].get(star, 0)
            s_ratio = round((s_cnt / len(reviews) * 100), 1) if len(reviews) > 0 else 0

            ws_dash.cell(row=c_row, column=1, value=label).alignment = align_center
            ws_dash.cell(row=c_row, column=2, value=s_cnt).alignment = align_right
            ws_dash.cell(row=c_row, column=3, value=f"{s_ratio}%").alignment = align_right
            for col_idx in range(1, 4):
                c = ws_dash.cell(row=c_row, column=col_idx)
                c.font = font_body
                c.border = thin_border
            c_row += 1

        # Table 3: Top Keywords
        kw_col_start = 6
        ws_dash.cell(row=4, column=kw_col_start, value="주요 키워드 빈도 TOP 15").font = font_bold
        headers_kw = ["순위", "키워드", "언급 빈도"]
        for idx, h in enumerate(headers_kw):
            cell = ws_dash.cell(row=5, column=kw_col_start + idx, value=h)
            cell.font = font_header
            cell.fill = fill_primary
            cell.alignment = align_center
            cell.border = thin_border

        kw_row = 6
        for rank, item in enumerate(analysis["top_keywords"], 1):
            ws_dash.cell(row=kw_row, column=kw_col_start, value=f"#{rank}").alignment = align_center
            ws_dash.cell(row=kw_row, column=kw_col_start + 1, value=item["word"]).alignment = align_center
            ws_dash.cell(row=kw_row, column=kw_col_start + 2, value=item["count"]).alignment = align_right
            for col_idx in range(kw_col_start, kw_col_start + 3):
                c = ws_dash.cell(row=kw_row, column=col_idx)
                c.font = font_body
                c.border = thin_border
            kw_row += 1

        # Auto-adjust summary widths
        for col in ws_dash.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_dash.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # ---------------- Helper for Review Sheets ----------------
        def write_reviews_to_sheet(sheet, items: List[Dict[str, Any]], sheet_title: str, header_fill):
            sheet.title = sheet_title
            sheet.views.sheetView[0].showGridLines = True

            sheet["A1"] = f"{sheet_title} (총 {len(items)}건)"
            sheet["A1"].font = font_title
            sheet["A2"] = f"추출 시각: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            sheet["A2"].font = font_sub

            headers = ["번호", "포털", "사업장명", "작성자", "평점/별점", "방문/작성일시", "인증유형", "리뷰 본문", "키워드/태그"]
            for col_idx, h in enumerate(headers, 1):
                cell = sheet.cell(row=4, column=col_idx, value=h)
                cell.font = font_header
                cell.fill = header_fill
                cell.alignment = align_center
                cell.border = thin_border

            for r_idx, r in enumerate(items, 1):
                row_num = 4 + r_idx
                sheet.cell(row=row_num, column=1, value=r_idx).alignment = align_center
                sheet.cell(row=row_num, column=2, value=r.get("portal", "")).alignment = align_center
                sheet.cell(row=row_num, column=3, value=r.get("place_name", "")).alignment = align_left
                sheet.cell(row=row_num, column=4, value=r.get("author", "")).alignment = align_left
                sheet.cell(row=row_num, column=5, value=r.get("rating") if r.get("rating") is not None else "-").alignment = align_center
                sheet.cell(row=row_num, column=6, value=r.get("date", "")).alignment = align_center
                sheet.cell(row=row_num, column=7, value=r.get("verification", "")).alignment = align_center
                sheet.cell(row=row_num, column=8, value=r.get("content", "")).alignment = align_left
                sheet.cell(row=row_num, column=9, value=r.get("tags", "")).alignment = align_left

                is_even = (r_idx % 2 == 0)
                for col_idx in range(1, 10):
                    c = sheet.cell(row=row_num, column=col_idx)
                    c.font = font_body
                    c.border = thin_border
                    if is_even:
                        c.fill = fill_zebra

            # Column widths
            widths = [8, 12, 22, 16, 12, 15, 14, 55, 25]
            for idx, w in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(idx)].width = w

        # Sheet 2: All Reviews
        ws_all = wb.create_sheet()
        write_reviews_to_sheet(ws_all, reviews, "📋 통합 전체 리뷰", fill_primary)

        # Sheet 3: Naver
        naver_reviews = [r for r in reviews if r.get("portal") == "Naver"]
        if naver_reviews:
            ws_naver = wb.create_sheet()
            write_reviews_to_sheet(ws_naver, naver_reviews, "🟢 네이버 플레이스", fill_naver)

        # Sheet 4: Kakao
        kakao_reviews = [r for r in reviews if r.get("portal") == "Kakao"]
        if kakao_reviews:
            ws_kakao = wb.create_sheet()
            write_reviews_to_sheet(ws_kakao, kakao_reviews, "🟡 카카오맵 (다음)", fill_primary)

        # Sheet 5: Google
        google_reviews = [r for r in reviews if r.get("portal") == "Google"]
        if google_reviews:
            ws_google = wb.create_sheet()
            write_reviews_to_sheet(ws_google, google_reviews, "🔵 구글 지도", fill_google)

        wb.save(filepath)
        return filepath

    @classmethod
    def export_csv(cls, reviews: List[Dict[str, Any]], filepath: str) -> str:
        """Exports review list into CSV format with utf-8-sig encoding for Excel compatibility."""
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
        headers = ["번호", "포털", "사업장명", "작성자", "평점/별점", "방문/작성일시", "인증유형", "리뷰 본문", "키워드/태그"]

        with open(filepath, mode="w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for idx, r in enumerate(reviews, 1):
                writer.writerow([
                    idx,
                    r.get("portal", ""),
                    r.get("place_name", ""),
                    r.get("author", ""),
                    r.get("rating", ""),
                    r.get("date", ""),
                    r.get("verification", ""),
                    r.get("content", ""),
                    r.get("tags", "")
                ])
        return filepath

    @classmethod
    def export_json(cls, reviews: List[Dict[str, Any]], filepath: str) -> str:
        """Exports review list into JSON format."""
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
        analysis = ReviewAnalyzer.analyze(reviews)
        data = {
            "exported_at": datetime.now().isoformat(),
            "analysis": analysis,
            "reviews": reviews
        }
        with open(filepath, mode="w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return filepath
